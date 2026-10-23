import io
import re
import docx
from docx.shared import Pt, Inches, RGBColor

def generate_resume_docx(cv_text: str, title: str = "Enhanced Resume") -> io.BytesIO:
    """
    Converts plain text or markdown-style enhanced resume into a
    professionally formatted Microsoft Word (.docx) document.
    """
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Document Header Title
    head_p = doc.add_paragraph()
    head_run = head_p.add_run(title.upper())
    head_run.font.size = Pt(16)
    head_run.bold = True
    head_run.font.color.rgb = RGBColor(124, 58, 237) # Brand Purple
    head_p.paragraph_format.space_after = Pt(12)

    lines = cv_text.splitlines()
    for line in lines:
        raw = line.strip()
        if not raw or raw.startswith('===') or raw == '***' or raw == '---':
            continue

        # Check for Section Headings e.g. [PROFESSIONAL SUMMARY], ### EXPERIENCE, etc.
        if (raw.startswith('[') and raw.endswith(']')) or raw.startswith('###') or raw.startswith('##'):
            clean_head = re.sub(r'^[#\s\[]+|[\]\s]+$', '', raw)
            p = doc.add_paragraph()
            r = p.add_run(clean_head)
            r.bold = True
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(30, 41, 59)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
        # Check for bullet points
        elif raw.startswith('•') or raw.startswith('* ') or raw.startswith('- '):
            clean_bullet = re.sub(r'^[•\*\-\s]+', '', raw)
            p = doc.add_paragraph(style='List Bullet')
            r = p.add_run(clean_bullet)
            r.font.size = Pt(10.5)
            p.paragraph_format.space_after = Pt(2)
        else:
            p = doc.add_paragraph()
            r = p.add_run(raw)
            r.font.size = Pt(10.5)
            p.paragraph_format.space_after = Pt(4)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
