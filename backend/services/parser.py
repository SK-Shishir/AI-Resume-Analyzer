import io
import os
from pypdf import PdfReader
from docx import Document

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF bytes."""
    reader = PdfReader(io.BytesIO(file_bytes))
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return "\n".join(text_parts).strip()

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX bytes."""
    doc = Document(io.BytesIO(file_bytes))
    text_parts = []
    
    # Extract from paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text.strip())
            
    # Extract from tables if any
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_parts.append(" | ".join(row_text))
                
    return "\n".join(text_parts).strip()

def parse_resume(file_bytes: bytes, filename: str) -> str:
    """Extract text based on file extension."""
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == '.pdf':
        return extract_text_from_pdf(file_bytes)
    elif ext == '.docx':
        return extract_text_from_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file format: {ext}. Only PDF and DOCX are supported.")
